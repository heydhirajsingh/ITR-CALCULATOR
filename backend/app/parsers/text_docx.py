"""TXT and DOCX parsers for statements, ledgers, certificates and salary records."""
from __future__ import annotations

import io
from pathlib import Path

import pandas as pd
from docx import Document as DocxDocument

from backend.app.parsers.base import DocumentParser, ParseResult, ParserError
from backend.app.parsers.detector import detect_bank, detect_document_type
from backend.app.parsers.pdf import _parse_text_lines
from backend.app.parsers.tabular import dataframe_to_transactions, rows_to_dataframe


class TextDocxParser(DocumentParser):
    extensions = (".txt", ".docx")

    def parse(self, path: Path, *, password: str | None = None) -> ParseResult:
        del password
        warnings: list[str] = []
        tables: list[pd.DataFrame] = []
        if path.suffix.lower() == ".docx":
            try:
                document = DocxDocument(path)
            except Exception as exc:
                raise ParserError(f"Could not read DOCX: {exc}") from exc
            text = "\n".join(paragraph.text for paragraph in document.paragraphs)
            for table in document.tables:
                rows = [[cell.text for cell in row.cells] for row in table.rows]
                frame = rows_to_dataframe(rows)
                if not frame.empty:
                    tables.append(frame)
        else:
            text = _read_text(path)
            try:
                frame = pd.read_csv(io.StringIO(text), sep=None, engine="python", dtype=object)
                if len(frame.columns) >= 3:
                    tables.append(frame)
            except Exception:
                pass

        document_type = detect_document_type(path.name, text)
        bank_name = detect_bank(path.name, text)
        transactions = []
        for frame in tables:
            transactions.extend(dataframe_to_transactions(frame, bank_name=bank_name))
        if not transactions:
            transactions = _parse_text_lines(text, bank_name, None)
        if not transactions:
            warnings.append("Text was extracted but no standard transaction rows were found.")
        return ParseResult(
            transactions=transactions,
            document_type=document_type,
            parser_used="python-docx/text",
            page_texts=[(1, text, "text", 100.0)],
            warnings=warnings,
            bank_name=bank_name,
            page_count=1,
        )


def _read_text(path: Path) -> str:
    for encoding in ("utf-8-sig", "utf-8", "cp1252", "latin-1"):
        try:
            return path.read_text(encoding=encoding)
        except UnicodeDecodeError:
            continue
    return path.read_text(errors="replace")
